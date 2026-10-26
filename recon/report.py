"""Reconciliation summary (bank-to-book) and CSV / Excel export."""
from __future__ import annotations

import csv
import io
from decimal import Decimal
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import ReconResult

ZERO = Decimal("0.00")


def effective_state(result: ReconResult, decisions: Optional[dict] = None):
    """Apply reviewer decisions: rejected matches and rejected duplicates go back
    to the unmatched lists."""
    decisions = decisions or {}
    rejected = {k for k, v in decisions.items() if v["decision"] == "reject"}
    matches = [m for m in result.matches if m.id not in rejected]
    unmatched_bank = list(result.unmatched_bank)
    unmatched_ledger = list(result.unmatched_ledger)
    for m in result.matches:
        if m.id in rejected:
            unmatched_bank += m.bank_ids
            unmatched_ledger += m.ledger_ids
    duplicates = []
    for d in result.duplicates:
        if d.id in rejected:
            (unmatched_bank if d.source == "bank" else unmatched_ledger).append(d.txn_id)
        else:
            duplicates.append(d)
    return matches, duplicates, unmatched_bank, unmatched_ledger


def build_summary(result: ReconResult, decisions: Optional[dict] = None,
                  opening_bank: Decimal = ZERO, opening_book: Decimal = ZERO) -> dict:
    matches, duplicates, ub, ul = effective_state(result, decisions)
    bank_bal = opening_bank + sum((t.amount for t in result.bank.values()), ZERO)
    book_bal = opening_book + sum((t.amount for t in result.ledger.values()), ZERO)

    in_transit = sum((result.ledger[i].amount for i in ul if result.ledger[i].amount > 0), ZERO)
    outstanding = sum((result.ledger[i].amount for i in ul if result.ledger[i].amount < 0), ZERO)
    bank_dups = sum((d.amount for d in duplicates if d.source == "bank"), ZERO)
    ledger_dups = sum((d.amount for d in duplicates if d.source == "ledger"), ZERO)
    bank_credits = sum((result.bank[i].amount for i in ub if result.bank[i].amount > 0), ZERO)
    bank_debits = sum((result.bank[i].amount for i in ub if result.bank[i].amount < 0), ZERO)
    diffs = sum((m.difference for m in matches), ZERO)

    adj_bank = bank_bal + in_transit + outstanding - bank_dups
    adj_book = book_bal + bank_credits + bank_debits + diffs - ledger_dups

    lines = [
        ("Bank side", "Balance per bank statement", bank_bal),
        ("Bank side", "Add: deposits in transit (in books, not yet in bank)", in_transit),
        ("Bank side", "Less: outstanding payments (in books, not yet in bank)", outstanding),
        ("Bank side", "Less: suspected duplicate bank lines (pending reversal)", -bank_dups),
        ("Bank side", "Adjusted bank balance", adj_bank),
        ("Book side", "Balance per ledger", book_bal),
        ("Book side", "Add: bank credits not recorded in books", bank_credits),
        ("Book side", "Less: bank debits not recorded in books", bank_debits),
        ("Book side", "Add/(less): amount differences on matched items (bank - book)", diffs),
        ("Book side", "Less: suspected duplicate ledger entries", -ledger_dups),
        ("Book side", "Adjusted book balance", adj_book),
        ("Result", "Unexplained difference", adj_bank - adj_book),
    ]
    return {
        "lines": [{"section": s, "label": l, "amount": f"{a:.2f}"} for s, l, a in lines],
        "balanced": adj_bank == adj_book,
        "counts": {
            "matches": len(matches),
            "duplicates": len(duplicates),
            "unmatched_bank": len(ub),
            "unmatched_ledger": len(ul),
        },
    }


def result_rows(result: ReconResult, decisions: Optional[dict] = None) -> list[dict]:
    """One flat row per bank/ledger line with its status - the main export."""
    decisions = decisions or {}
    matches, duplicates, ub, ul = effective_state(result, decisions)
    rows = []

    def row(t, status, match=None, extra=""):
        c = result.classifications.get(t.id) if t.source == "bank" else None
        dec = decisions.get(match.id) if match else None
        return {
            "status": status,
            "match_id": match.id if match else "",
            "match_type": match.type if match else "",
            "confidence": f"{match.confidence:.3f}" if match else "",
            "review": (dec["decision"] if dec else ("pending" if match and match.needs_review else "")),
            "source": t.source,
            "txn_id": t.id,
            "date": t.date.isoformat(),
            "description": t.description,
            "amount": f"{t.amount:.2f}",
            "reference": t.reference,
            "difference": f"{match.difference:.2f}" if match else "",
            "suggested_account": c.account if c else "",
            "account_source": c.source if c else "",
            "notes": extra or (match.reason if match else ""),
        }

    for m in matches:
        status = "amount_difference" if m.type == "amount_difference" else "matched"
        for b in m.bank_ids:
            rows.append(row(result.bank[b], status, m))
        for l in m.ledger_ids:
            rows.append(row(result.ledger[l], status, m))
    for d in duplicates:
        t = (result.bank if d.source == "bank" else result.ledger)[d.txn_id]
        rows.append(row(t, "suspected_duplicate", extra=d.reason))
    for b in ub:
        rows.append(row(result.bank[b], "unmatched"))
    for l in ul:
        rows.append(row(result.ledger[l], "unmatched"))
    return rows


def _to_csv(rows: list[dict]) -> str:
    buf = io.StringIO()
    if rows:
        w = csv.DictWriter(buf, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)
    return buf.getvalue()


def export_results_csv(result, decisions=None) -> str:
    return _to_csv(result_rows(result, decisions))


def export_summary_csv(result, decisions=None) -> str:
    return _to_csv(build_summary(result, decisions)["lines"])


def _write_sheet(ws, rows: list[dict], money_cols=("amount", "difference")):
    if not rows:
        ws.append(["(none)"])
        return
    headers = list(rows[0].keys())
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
    for r in rows:
        values = []
        for h in headers:
            v = r[h]
            if h in money_cols and v not in ("", None):
                v = float(v)
            elif isinstance(v, (dict, list)):
                v = str(v)
            values.append(v)
        ws.append(values)
    for idx, h in enumerate(headers, start=1):
        width = max(len(str(h)), *(len(str(r[h])) for r in rows)) + 2
        ws.column_dimensions[get_column_letter(idx)].width = min(width, 60)
        if h in money_cols:
            for cell in ws[get_column_letter(idx)][1:]:
                cell.number_format = "#,##0.00"
    ws.freeze_panes = "A2"


def export_xlsx(result, decisions=None, review_items=None, audit_entries=None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    summary = build_summary(result, decisions)
    _write_sheet(ws, summary["lines"])
    for row in ws.iter_rows(min_row=2):
        if row[1].value and ("Adjusted" in row[1].value or "Unexplained" in row[1].value):
            for c in row:
                c.font = Font(bold=True)
    rows = result_rows(result, decisions)
    for title, status in [("Matched", "matched"), ("Amount Differences", "amount_difference"),
                          ("Duplicates", "suspected_duplicate"), ("Unmatched", "unmatched")]:
        _write_sheet(wb.create_sheet(title), [r for r in rows if r["status"] == status])
    if review_items is not None:
        _write_sheet(wb.create_sheet("Review Queue"), review_items)
    if audit_entries is not None:
        _write_sheet(wb.create_sheet("Audit Log"), audit_entries)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
