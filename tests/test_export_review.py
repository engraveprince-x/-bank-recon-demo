import csv
import io
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from recon.pipeline import run
from recon.report import build_summary, export_results_csv, export_summary_csv, export_xlsx
from recon.review import AuditLog, ReviewSession

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture()
def result(monkeypatch):
    monkeypatch.delenv("LLM_API_KEY", raising=False)
    return run(ROOT / "data/sample/bank_statement.csv", ROOT / "data/sample/ledger.csv")


def test_summary_balances(result):
    s = build_summary(result)
    assert s["balanced"]
    assert s["lines"][-1]["amount"] == "0.00"


def test_results_csv_covers_every_row(result):
    rows = list(csv.DictReader(io.StringIO(export_results_csv(result))))
    ids = {(r["source"], r["txn_id"]) for r in rows}
    assert len(ids) == len(rows) == len(result.bank) + len(result.ledger)
    statuses = {r["status"] for r in rows}
    assert statuses == {"matched", "amount_difference", "suspected_duplicate", "unmatched"}
    unmatched_bank = [r for r in rows if r["status"] == "unmatched" and r["source"] == "bank"]
    assert all(r["suggested_account"] and r["account_source"] == "rule" for r in unmatched_bank)


def test_summary_csv(result):
    rows = list(csv.DictReader(io.StringIO(export_summary_csv(result))))
    assert rows[0]["label"] == "Balance per bank statement"
    assert rows[-1]["label"] == "Unexplained difference"


def test_xlsx_export(result, tmp_path):
    session = ReviewSession(result, AuditLog(tmp_path / "audit.jsonl"))
    session.decide(result.duplicates[0].id, "confirm", "alice")
    data = export_xlsx(result, session.decisions, session.items(), session.audit.entries())
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Summary", "Matched", "Amount Differences", "Duplicates", "Unmatched", "Review Queue", "Audit Log"]
    assert wb["Summary"]["B2"].value == "Balance per bank statement"
    assert wb["Amount Differences"].max_row == 1 + 2 * 4  # header + bank/ledger line per difference
    assert wb["Audit Log"]["B2"].value == "alice"


def test_review_reject_moves_items_back_and_is_audited(result, tmp_path):
    log = tmp_path / "audit.jsonl"
    session = ReviewSession(result, AuditLog(log))
    m = next(m for m in result.matches if m.needs_review and m.type == "fuzzy")
    before = build_summary(result, session.decisions)["counts"]
    session.decide(m.id, "reject", "bob", "wrong vendor")
    after = build_summary(result, session.decisions)
    assert after["counts"]["matches"] == before["matches"] - 1
    assert after["counts"]["unmatched_bank"] == before["unmatched_bank"] + 1
    assert after["balanced"]
    entry = AuditLog(log).entries()[-1]  # persisted to disk
    assert entry["user"] == "bob" and entry["action"] == "review.reject" and entry["target"] == m.id and entry["at"]
    with pytest.raises(ValueError):
        session.decide(m.id, "maybe", "bob")
    with pytest.raises(ValueError):
        session.decide(m.id, "confirm", "  ")
    with pytest.raises(KeyError):
        session.decide("M9999", "confirm", "bob")


def test_api_end_to_end(tmp_path, monkeypatch):
    monkeypatch.delenv("LLM_API_KEY", raising=False)
    from app.main import create_app
    client = TestClient(create_app(audit_path=tmp_path / "audit.jsonl"))
    assert client.get("/api/state").json()["loaded"] is False
    assert client.get("/api/export/results.csv").status_code == 409
    s = client.post("/api/sample", data={"user": "carol"}).json()
    assert s["loaded"] and s["summary"]["balanced"]
    item = s["review"][0]["item_id"]
    s = client.post(f"/api/review/{item}", json={"decision": "confirm", "user": "carol"}).json()
    assert next(i for i in s["review"] if i["item_id"] == item)["status"] == "confirm"
    assert s["audit"][0]["action"] == "review.confirm"
    assert client.post(f"/api/review/{item}", json={"decision": "confirm", "user": ""}).status_code == 400
    assert client.get("/api/export/report.xlsx").content[:2] == b"PK"
    bank = (ROOT / "data/sample/bank_statement.csv").read_bytes()
    bad = client.post("/api/upload", files={"bank": ("b.csv", b"foo,bar\n1,2\n"), "ledger": ("l.csv", bank)})
    assert bad.status_code == 400
    assert client.get("/").status_code == 200
